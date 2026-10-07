"""Render the direction mirror into a simple DOCX and verify text parity.

Run with the Codex artifact runtime's Python (python-docx), not the collector venv.
Layout inspection is a separate step; this tool does not claim visual validation.
"""
from pathlib import Path
import os
import tempfile
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]


def entries():
    for line in (ROOT / 'PROJECT_DIRECTION.md').read_text(encoding='utf-8').splitlines():
        if not line.strip():
            continue
        if line.startswith('# '):
            yield 'Title', line[2:]
        elif line.startswith('## '):
            yield 'Heading 1', line[3:]
        elif line.startswith('- '):
            yield 'List Bullet', line[2:]
        else:
            yield 'Normal', line


def main():
    document = Document()
    document.core_properties.title = 'Stock Helper 프로젝트 방향 기준서'
    document.core_properties.author = ''
    document.core_properties.last_modified_by = ''
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.2)
    for name, size in [('Normal',10.5),('Title',20),('Heading 1',13),('List Bullet',10.5)]:
        style = document.styles[name]
        style.font.name = 'Malgun Gothic'
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0,0,0)
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Malgun Gothic')
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.widow_control = True
    for style, text in entries():
        document.add_paragraph(text, style)
    handle, temporary = tempfile.mkstemp(suffix='.docx', dir=ROOT)
    os.close(handle)
    try:
        document.save(temporary)
        opened = Document(temporary)
        assert [p.text for p in opened.paragraphs] == [t for _,t in entries()]
        os.replace(temporary, ROOT/'PROJECT_DIRECTION.docx')
    finally:
        Path(temporary).unlink(missing_ok=True)
    print('DOCX open/read and all paragraph text parity: OK')


if __name__ == '__main__':
    main()
